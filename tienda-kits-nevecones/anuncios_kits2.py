#!/usr/bin/env python3
"""Versión compacta: 4 grupos de anuncios (en vez de 9). Las palabras clave que van a otra página llevan su propia URL final."""
import json, sys, html, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
OUT = sys.argv[1]
B = "https://filtrosaguanevecones.com/kits-instalacion-nevecones/"
G = json.load(open(OUT + "/datos.json"))["grupos"]; NEG = json.load(open(OUT + "/datos.json"))["neg"]
k = {g["n"][:2]: g["k"] for g in G}
P = lambda ks: [x for q in ks for x in (f'"{q}"', f'[{q}]')]
GR = [
 dict(n="A Kits completos con filtro", url=B, path=("kit-completo", "con-filtro"),
  t=["Kit Agua para tu Nevecón","Manguera + Llave + Filtro","Kit Completo con Filtro","Conecta el Agua de tu Nevecón","Filtro Externo Incluido","Todo en un Solo Pedido","Kit Completo Desde $89.900","Ahorra Comprando el Kit","Manguera 5 m y Llave Galápago","Envío Gratis a Toda Colombia","Compra Segura en Línea","Paga con PSE o Nequi"],
  d=["Kit completo: manguera 1/4\" de 5 m, llave galápago y filtro externo. Envío gratis.","Todo para conectar el agua de tu nevecón en un solo pedido. Paga con PSE o Nequi.","Ahorra más que comprando las piezas por separado. Llega en 2 a 3 días hábiles.","Con paso a paso para instalarlo tú mismo. Pide tu kit hoy y recíbelo en casa."],
  kb=[("Samsung", B+"kit-completo-nevecon-samsung/", P(k["01"])), ("LG", B+"kit-completo-nevecon-lg/", P(k["02"])), ("Whirlpool", B+"kit-completo-nevecon-whirlpool/", P(k["03"]))]),
 dict(n="B Kit básico (manguera + llave)", url=B+"kit-basico-manguera-llave-galapago/", path=("kit-nevecon", "basico"),
  t=[x for x in G[3]["t"]], d=G[3]["d"], kb=[("Todas", B+"kit-basico-manguera-llave-galapago/", P(k["04"]))]),
 dict(n="C Mangueras y acoples", url=B+"manguera-1-4-para-nevecon/", path=("manguera", "nevecon"),
  t=["Manguera para Nevecón","Manguera 1/4\" Desde $38.900","Manguera 5/16\" Desde $41.900","Manguera 5, 10, 15 o 20 m","Acople Rápido Desde $29.900","Une Mangueras sin Herramientas","Escoge los Metros que Quieras","Manguera de Agua para Nevera","Envío Gratis a Toda Colombia","Compra Segura en Línea","Paga con PSE o Nequi","Llega en 2 a 3 Días"],
  d=["Manguera de 1/4\" o 5/16\" de 5 a 20 metros y acoples rápidos. Envío gratis a toda Colombia.","Escoge el largo que necesitas para conectar el agua de tu nevecón.","Paga con tarjeta, PSE o Nequi. Recíbelo en 2 a 3 días hábiles.","Combínalo con llave y filtro en un kit y ahorra. Compra segura en línea."],
  kb=[("Manguera 1/4\"", B+"manguera-1-4-para-nevecon/", P(k["06"])), ("Manguera 5/16\"", B+"manguera-5-16-para-nevecon/", P(k["07"])), ("Acoples rápidos", B+"acople-rapido-manguera-1-4/", P(k["08"]))]),
 dict(n="D Llaves para conectar el nevecón", url=B+"llave-galapago-para-nevecon/", path=("llave", "nevecon"),
  t=["Llave Galápago para Nevecón","Llave Galápago Desde $48.900","Llave Plástica Desde $53.900","Con Racores, Anillo y Tuerca","Toma Agua sin Cortar el Tubo","Llave para Conectar Nevecón","Llave con Acople Rápido","Envío Gratis a Toda Colombia","Compra Segura en Línea","Paga con PSE o Nequi","Fácil de Instalar","Llega en 2 a 3 Días"],
  d=["Llave galápago o llave plástica con acople rápido para conectar el agua de tu nevecón.","Toma el agua del tubo sin cortarlo. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Incluye paso a paso para instalarla.","Pídela hoy y recíbela en 2 a 3 días hábiles. Compra segura en línea."],
  kb=[("Llave galápago", B+"llave-galapago-para-nevecon/", P(k["05"])), ("Llave plástica", B+"llave-plastica-acople-rapido-para-nevecon/", P(k["09"]))]),
]
for g in GR:
    assert all(len(t) <= 30 for t in g["t"]), [t for t in g["t"] if len(t) > 30]
    assert all(len(d) <= 90 for d in g["d"]), [(len(d), d) for d in g["d"] if len(d) > 90]
    assert len(set(g["t"])) == len(g["t"])
e = html.escape
def lista(i, items, tit, ctr=False):
    li = "".join(f'<li><span class="tx">{e(x)}</span>{f"<small>{len(x)}</small>" if ctr else ""}<button type="button" class="c1">Copiar</button></li>' for x in items)
    return f'<div class="bloque"><div class="bh"><b>{tit}</b><button type="button" class="ca" data-t="{i}">Copiar todos</button></div><ul id="{i}">{li}</ul></div>'
sec = ""
for i, g in enumerate(GR, 1):
    kbs = ""
    for j, (et, u, ks) in enumerate(g["kb"]):
        kbs += (f'<div class="bloque sub"><div class="bh"><b>Palabras clave: {e(et)}</b></div>'
                f'<div class="bh"><span class="et">URL final de estas palabras</span><button type="button" class="c1">Copiar</button></div><p class="tx url">{e(u)}</p>' if len(g["kb"]) > 1 else f'<div class="bloque sub">') \
               + lista(f"k{i}_{j}", ks, "Lista (frase y exacta)") + "</div>"
    sec += f'''<article class="g" id="g{i}"><h2>Grupo {e(g['n'])}</h2>
<div class="bloque"><div class="bh"><b>URL final del anuncio</b><button type="button" class="c1">Copiar</button></div><p class="tx url">{e(g['url'])}</p></div>
<div class="bloque"><div class="bh"><b>Ruta visible</b></div><p><span class="tx">{e(g['path'][0])}</span> <button type="button" class="c1">Copiar</button> / <span class="tx">{e(g['path'][1])}</span> <button type="button" class="c1">Copiar</button></p></div>
{lista(f't{i}', g['t'], 'Titulares (máx. 30 letras; el número es su largo)', True)}
{lista(f'd{i}', g['d'], 'Descripciones (máx. 90 letras)', True)}
<div class="bloque"><div class="bh"><b>Palabras clave</b></div>{kbs}</div></article>'''
sec += f'<article class="g" id="neg-s"><h2>Palabras negativas de la campaña</h2>{lista("neg", NEG, "Palabras clave negativas (frase)")}</article>'
tpl = open(OUT + "/index.html").read()
cab = tpl[:tpl.index('<div class="w">')]; pie = tpl[tpl.index("<script>"):]
cab = cab.replace("Anuncios Kits Google Ads", "Anuncios Kits 4 Grupos").replace(".bloque{", ".bloque.sub{padding:8px;border:1px dashed var(--line);border-radius:10px}.bloque{")
nav = "".join(f'<a href="#g{i}">{n}</a>' for i, n in enumerate("ABCD", 1)) + '<a href="#neg-s">Negativas</a>'
page = cab + f'''<div class="w"><h1>Anuncios de Google Ads para los kits: 4 grupos</h1>
<div class="aviso"><b>Cómo se usa:</b> crea UNA campaña de búsqueda y dentro de ella estos 4 grupos de anuncios. En cada grupo copia los titulares, las descripciones y las palabras clave con su botón. <b>Ojo:</b> en los grupos con varias páginas, cada lista de palabras lleva su propia «URL final»: en Google Ads se pone a nivel de palabra clave (columna «URL final»). Las negativas van una sola vez, en la campaña.</div>
<nav class="ix">{nav}</nav>{sec}</div>''' + pie
open(OUT + "/index2.html", "w").write(page)
wb = openpyxl.Workbook(); wb.remove(wb.active)
def hoja(n, cab, filas, an):
    ws = wb.create_sheet(n); ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1236A8")
    for f in filas: ws.append(f)
    for i, w in enumerate(an): ws.column_dimensions[chr(65 + i)].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
hoja("Grupos", ["Campaña", "Grupo", "URL final", "Ruta 1", "Ruta 2"], [["Kits instalación nevecón - Colombia", g["n"], g["url"], g["path"][0], g["path"][1]] for g in GR], [38, 36, 80, 16, 16])
hoja("Titulares", ["Grupo", "Titular (máx 30)", "Largo"], [[g["n"], t, len(t)] for g in GR for t in g["t"]], [36, 36, 8])
hoja("Descripciones", ["Grupo", "Descripción (máx 90)", "Largo"], [[g["n"], d, len(d)] for g in GR for d in g["d"]], [36, 96, 8])
hoja("Palabras clave", ["Grupo", "Palabra clave", "Concordancia", "URL final de la palabra"], [[g["n"], x, "Exacta" if x.startswith("[") else "Frase", u] for g in GR for (_e, u, ks) in g["kb"] for x in ks], [36, 52, 14, 80])
hoja("Negativas (campaña)", ["Palabra clave negativa", "Tipo"], [[n, "Frase"] for n in NEG], [34, 12])
wb.save(OUT + "/anuncios-kits-4-grupos.xlsx")
print("ok", sum(len(ks) for g in GR for (_e, _u, ks) in g["kb"]), "palabras clave")
