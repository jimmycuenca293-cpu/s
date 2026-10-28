#!/usr/bin/env python3
"""Grupos de anuncios de Google Ads para los productos nuevos (kits de instalación). Valida los largos: títulos <=30, descripciones <=90."""
import sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
OUT = sys.argv[1]
BASE = "https://filtrosaguanevecones.com/kits-instalacion-nevecones/"
COMUN_T = ["Envío Gratis a Toda Colombia", "Paga con Tarjeta, PSE o Nequi", "Compra Segura en Línea", "Te Llega en 2 a 3 Días Hábiles"]
GRUPOS = [
 dict(n="01 Kit completo Samsung", url=BASE+"kit-completo-nevecon-samsung/", path=("kit-nevecon", "samsung"),
  t=["Kit Agua para Nevecón Samsung","Manguera + Llave + Filtro","Kit Completo Nevecón Samsung","Conecta el Agua de tu Nevecón","Kit con Filtro de tu Marca","Todo en un Solo Pedido","Kit Instalación Desde $89.900","Ahorra Comprando el Kit","Manguera 5 m y Llave Galápago","Incluye Filtro Externo","Envío Gratis a Toda Colombia","Compra Segura en Línea"],
  d=["Kit completo: manguera 1/4\" de 5 m, llave galápago y filtro externo. Envío gratis.","Todo para conectar el agua de tu nevecón en un solo pedido. Paga con PSE o Nequi.","Ahorra más que comprando las piezas por separado. Llega en 2 a 3 días hábiles.","Con paso a paso para instalarlo tú mismo. Pide tu kit hoy y recíbelo en casa."],
  k=["kit instalación nevecón samsung","kit agua nevecón samsung","kit conexión agua nevecón samsung","conectar agua nevecón samsung","manguera y llave para nevecón samsung","kit instalación nevera samsung","kit filtro externo nevecón samsung","instalar agua nevecón samsung"]),
 dict(n="02 Kit completo LG", url=BASE+"kit-completo-nevecon-lg/", path=("kit-nevecon", "lg"),
  t=["Kit Agua para Nevecón LG","Manguera + Llave + Filtro","Kit Completo Nevecón LG","Conecta el Agua de tu Nevecón","Kit con Filtro de tu Marca","Todo en un Solo Pedido","Kit Instalación Desde $89.900","Ahorra Comprando el Kit","Manguera 5 m y Llave Galápago","Incluye Filtro Externo","Envío Gratis a Toda Colombia","Compra Segura en Línea"],
  d=["Kit completo: manguera 1/4\" de 5 m, llave galápago y filtro externo. Envío gratis.","Todo para conectar el agua de tu nevecón en un solo pedido. Paga con PSE o Nequi.","Ahorra más que comprando las piezas por separado. Llega en 2 a 3 días hábiles.","Con paso a paso para instalarlo tú mismo. Pide tu kit hoy y recíbelo en casa."],
  k=["kit instalación nevecón lg","kit agua nevecón lg","kit conexión agua nevecón lg","conectar agua nevecón lg","manguera y llave para nevecón lg","kit instalación nevera lg","kit filtro externo nevecón lg","instalar agua nevecón lg"]),
 dict(n="03 Kit completo Whirlpool", url=BASE+"kit-completo-nevecon-whirlpool/", path=("kit-nevecon", "whirlpool"),
  t=["Kit Agua Nevecón Whirlpool","Manguera + Llave + Filtro","Kit Completo Whirlpool","Conecta el Agua de tu Nevecón","Kit con Filtro de tu Marca","Todo en un Solo Pedido","Kit Instalación Desde $97.900","Ahorra Comprando el Kit","Manguera 5 m y Llave Galápago","Incluye Filtro Externo","Envío Gratis a Toda Colombia","Compra Segura en Línea"],
  d=["Kit completo: manguera 1/4\" de 5 m, llave galápago y filtro externo. Envío gratis.","Todo para conectar el agua de tu nevecón en un solo pedido. Paga con PSE o Nequi.","Ahorra más que comprando las piezas por separado. Llega en 2 a 3 días hábiles.","Con paso a paso para instalarlo tú mismo. Pide tu kit hoy y recíbelo en casa."],
  k=["kit instalación nevecón whirlpool","kit agua nevecón whirlpool","kit conexión agua nevecón whirlpool","conectar agua nevecón whirlpool","manguera y llave para nevecón whirlpool","kit instalación nevera whirlpool","kit filtro externo nevecón whirlpool","instalar agua nevecón whirlpool"]),
 dict(n="04 Kit básico (manguera + llave)", url=BASE+"kit-basico-manguera-llave-galapago/", path=("kit-nevecon", "basico"),
  t=["Kit Instalación para Nevecón","Manguera 1/4\" + Llave Galápago","Kit Básico Desde $66.900","Conecta el Agua del Nevecón","Manguera 5 Metros Incluida","Llave Galápago con Racores","Envío Gratis a Toda Colombia","Compra Segura en Línea","Listo para Instalar","Kit Agua para Nevera","Con Anillo y Tuerca","Paga con PSE o Nequi"],
  d=["Kit con manguera de 1/4\" de 5 metros y llave galápago con racores, anillo y tuerca.","Lo mínimo para conectar el agua de tu nevecón. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Recíbelo en 2 a 3 días hábiles.","Incluye paso a paso para instalarlo. Pide tu kit básico hoy."],
  k=["kit instalación nevecón","kit de instalación para nevecón","kit instalación nevera","kit agua nevecón","kit conexión agua refrigerador","kit manguera y llave nevecón","manguera y llave galápago nevecón","kit instalación agua nevera"]),
 dict(n="05 Llave galápago", url=BASE+"llave-galapago-para-nevecon/", path=("llave", "galapago"),
  t=["Llave Galápago para Nevecón","Llave Galápago Desde $48.900","Con Racores, Anillo y Tuerca","Toma Agua sin Cortar el Tubo","Llave para Conectar Nevecón","Envío Gratis a Toda Colombia","Compra Segura en Línea","Llave Galápago para Nevera","Fácil de Instalar","Paga con PSE o Nequi","Para Manguera de 1/4\"","Llega en 2 a 3 Días"],
  d=["Llave galápago con racores, anillo y tuerca para conectar la manguera de tu nevecón.","Toma el agua del tubo sin cortarlo. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Incluye paso a paso para instalarla.","Pídela hoy y recíbela en 2 a 3 días hábiles. Compra segura en línea."],
  k=["llave galápago","llave galapago","llave galápago para nevera","llave galápago nevecón","llave galápago 1/4","llave galápago agua nevera","comprar llave galápago","llave para conectar nevecón al agua","llave galápago con racores"]),
 dict(n="06 Manguera 1/4\" para nevecón", url=BASE+"manguera-1-4-para-nevecon/", path=("manguera", "1-4"),
  t=["Manguera 1/4\" para Nevecón","Manguera Desde $38.900","Manguera 5, 10, 15 o 20 m","Conecta el Agua del Nevecón","Manguera para Dispensador","Envío Gratis a Toda Colombia","Compra Segura en Línea","Manguera de Agua para Nevera","Escoge los Metros que Quieras","Paga con PSE o Nequi","Manguera Blanca 1/4 Pulgada","Llega en 2 a 3 Días"],
  d=["Manguera de 1/4\" para conectar el agua de tu nevecón: 5, 10, 15 o 20 metros.","Escoge el largo que necesitas. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Recíbela en 2 a 3 días hábiles.","Combínala con llave y acoples en un kit y ahorra. Compra segura en línea."],
  k=["manguera 1/4 nevecón","manguera para nevecón","manguera de agua para nevera","manguera 1/4 para nevera","manguera dispensador agua nevecón","manguera 1/4 5 metros","manguera nevecón 10 metros","manguera agua refrigerador","comprar manguera nevecón"]),
 dict(n="07 Manguera 5/16\"", url=BASE+"manguera-5-16-para-nevecon/", path=("manguera", "5-16"),
  t=["Manguera 5/16\" para Nevecón","Manguera 5/16 Desde $41.900","Manguera 5, 10, 15 o 20 m","Manguera Transparente 5/16","Envío Gratis a Toda Colombia","Compra Segura en Línea","Escoge los Metros que Quieras","Paga con PSE o Nequi","Manguera de Agua 5/16 Pulgada","Llega en 2 a 3 Días","Para Conexión de Agua","Pide tu Manguera Hoy"],
  d=["Manguera transparente de 5/16\" para conexiones de agua: 5, 10, 15 o 20 metros.","Escoge el largo que necesitas. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Recíbela en 2 a 3 días hábiles.","¿No sabes la medida? Escríbenos por WhatsApp con una foto y te ayudamos."],
  k=["manguera 5/16","manguera 5/16 agua","manguera 5/16 nevera","manguera 5/16 pulgada","manguera transparente 5/16","manguera 5/16 5 metros","manguera 5/16 10 metros","comprar manguera 5/16"]),
 dict(n="08 Acoples rápidos", url=BASE+"acople-rapido-manguera-1-4/", path=("acople", "rapido"),
  t=["Acople Rápido Manguera 1/4\"","Acople Rápido 1/4 y 5/16","Acople Rápido Desde $29.900","Une Mangueras sin Herramientas","Empalme Rápido para Manguera","Envío Gratis a Toda Colombia","Compra Segura en Línea","Acople Recto para Agua","Paga con PSE o Nequi","Llega en 2 a 3 Días","Alarga tu Manguera de Agua","Pide tu Acople Hoy"],
  d=["Acople rápido de unión recta para empalmar manguera de 1/4\" o 5/16\", sin herramientas.","Empuja la manguera y queda sujeta. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Recíbelo en 2 a 3 días hábiles.","Úsalo para alargar o reparar la manguera de agua de tu nevecón."],
  k=["acople rápido manguera 1/4","acople rápido 1/4","acople rápido 5/16","acople rápido para manguera de agua","unión rápida manguera 1/4","empalme rápido manguera agua","acople rápido nevecón","comprar acople rápido 1/4"]),
 dict(n="09 Llave plástica con acople rápido", url=BASE+"llave-plastica-acople-rapido-para-nevecon/", path=("llave", "plastica"),
  t=["Llave Plástica para Nevecón","Llave con Acople Rápido","Llave Plástica Desde $53.900","Conecta sin Herramientas","Envío Gratis a Toda Colombia","Compra Segura en Línea","Paga con PSE o Nequi","Fácil de Instalar","Llega en 2 a 3 Días","Para Conectar el Agua","Con Tuerca y Acople Azul","Pide tu Llave Hoy"],
  d=["Llave plástica con acople rápido y tuerca para conectar la manguera de tu nevecón.","Empuja la manguera y queda sujeta. Envío gratis a toda Colombia.","Paga con tarjeta, PSE o Nequi. Recíbela en 2 a 3 días hábiles.","Incluye paso a paso. ¿Dudas de la rosca? Escríbenos por WhatsApp."],
  k=["llave plástica nevecón","llave con acople rápido nevera","llave plástica acople rápido","adaptador manguera nevecón","llave para conectar nevera al agua","llave plástica para nevera"]),
]
NEG = ["gratis","pdf","manual","tutorial","curso","empleo","trabajo","usado","segunda","olx","mercadolibre","homecenter","falabella","alkosto","cobre","gas","refrigerante","compresor","aire acondicionado","aire comprimido","neumática","piscina","jardín","riego","desagüe","lavadora","lavaplatos","grifo","técnico","servicio técnico","reparación","repuesto lavadora","bomba","filtro interno","ice maker","hielo"]
bad = []
for g in GRUPOS:
    for t in g["t"]:
        if len(t) > 30: bad.append(("titular", g["n"], len(t), t))
    for t in COMUN_T:
        if len(t) > 30: bad.append(("titular común", len(t), t))
    for d in g["d"]:
        if len(d) > 90: bad.append(("descripción", g["n"], len(d), d))
    assert len(g["t"]) == len(set(g["t"])), g["n"]
print("errores de largo:", bad)
wb = openpyxl.Workbook()
def hoja(nombre, cab, filas, anchos):
    ws = wb.create_sheet(nombre); ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1236A8")
    for f in filas: ws.append(f)
    for i, w in enumerate(anchos): ws.column_dimensions[chr(65 + i)].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
wb.remove(wb.active)
hoja("Grupos de anuncios", ["Campaña", "Grupo de anuncios", "URL final", "Ruta 1 (máx 15)", "Ruta 2 (máx 15)"], [["Kits instalación nevecón - Colombia", g["n"], g["url"], g["path"][0], g["path"][1]] for g in GRUPOS], [38, 36, 80, 16, 16])
hoja("Titulares", ["Grupo de anuncios", "Titular (máx 30)", "Largo"], [[g["n"], t, len(t)] for g in GRUPOS for t in g["t"]], [36, 36, 8])
hoja("Descripciones", ["Grupo de anuncios", "Descripción (máx 90)", "Largo"], [[g["n"], d, len(d)] for g in GRUPOS for d in g["d"]], [36, 96, 8])
filas = []
for g in GRUPOS:
    for k in g["k"]: filas += [[g["n"], f'"{k}"', "Frase"], [g["n"], f"[{k}]", "Exacta"]]
hoja("Palabras clave", ["Grupo de anuncios", "Palabra clave", "Tipo de concordancia"], filas, [36, 52, 20])
hoja("Negativas (campaña)", ["Palabra clave negativa", "Tipo"], [[n, "Frase"] for n in NEG], [34, 12])
wb.save(OUT + "/anuncios-kits-google-ads.xlsx")
import json; json.dump({"grupos": GRUPOS, "neg": NEG}, open(OUT + "/datos.json", "w"), ensure_ascii=False)
print(len(GRUPOS), "grupos;", sum(len(g["k"]) for g in GRUPOS), "palabras clave;", sum(len(g["t"]) for g in GRUPOS), "titulares")
